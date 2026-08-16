#!/usr/bin/env python3
"""Rebuilds the Character/Technique Templates spreadsheet from data/*.json.

Normally you go the other way (edit the spreadsheet, run
spreadsheet_to_json.py) — this script exists for the reverse case: the JSON
changed some other way and the spreadsheet needs to catch up, or the
spreadsheet file itself got lost/corrupted and needs regenerating from the
last-known-good JSON.

Usage:  python tools/json_to_spreadsheet.py
Reads:  data/characters.json, data/techniques.json
Writes: data/character-templates.xlsx
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, 'data')
OUT_PATH = os.path.join(DATA_DIR, 'character-templates.xlsx')

HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='4B3B2A')
BODY_FONT = Font(name='Arial', size=11)
NOTE_FONT = Font(name='Arial', size=11, italic=True, color='555555')

CHAR_COLUMNS = [
    ('key', 'Key'), ('archetype', 'Archetype'), ('category', 'Category'), ('species', 'Species'),
    ('woundsMax', 'Wounds Max (y)'), ('ap', 'AP'), ('flowMax', 'Flow'), ('recovery', 'Recovery'),
    ('power', 'Power'), ('quickness', 'Quickness'), ('resilience', 'Resilience'),
    ('lightning', 'Lightning'), ('lava', 'Lava'), ('focus', 'Focus'), ('control', 'Control'),
    ('attunement', 'Attunement'), ('life', 'Life'),
    ('techFlowPowers', 'Tech: Flow Powers'), ('techFocusPowers', 'Tech: Focus Powers'),
    ('techLightningCrafts', 'Tech: Lightning Crafts'), ('techLavaCrafts', 'Tech: Lava Crafts'),
]

TECH_COLUMNS = [
    ('key', 'Key'), ('name', 'Name'), ('group', 'Group'), ('column', 'Column'), ('type', 'Type'),
    ('apCost', 'AP Cost'), ('focusCost', 'Focus Cost'), ('flowCost', 'Flow Cost'),
    ('lightningCost', 'Lightning Cost'), ('lavaCost', 'Lava Cost'), ('range', 'Range'),
    ('durationSequences', 'Duration (Sequences)'), ('targetSide', 'Target Side'), ('effect', 'Effect'),
]


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30


def write_characters(wb, characters):
    ws = wb.active
    ws.title = 'Characters'
    for i, (_, label) in enumerate(CHAR_COLUMNS, start=1):
        ws.cell(row=1, column=i, value=label)
    style_header(ws, len(CHAR_COLUMNS))

    r = 2
    for key, row in characters.items():
        access = row.get('techniqueAccess') or {}
        values = {
            'key': key, 'archetype': row.get('archetype', ''), 'category': row.get('category', 'enemy'),
            'species': row.get('species') or '',
            'woundsMax': row.get('woundsMax'), 'ap': row.get('ap'), 'flowMax': row.get('flowMax'),
            'recovery': row.get('recovery'), 'power': row.get('power'), 'quickness': row.get('quickness'),
            'resilience': row.get('resilience'), 'lightning': row.get('lightning'), 'lava': row.get('lava'),
            'focus': row.get('focus'), 'control': row.get('control'), 'attunement': row.get('attunement'),
            'life': row.get('life', 1),
            'techFlowPowers': access.get('flowPowers', True), 'techFocusPowers': access.get('focusPowers', True),
            'techLightningCrafts': access.get('lightningCrafts', True), 'techLavaCrafts': access.get('lavaCrafts', True),
        }
        for i, (field, _) in enumerate(CHAR_COLUMNS, start=1):
            cell = ws.cell(row=r, column=i, value=values[field])
            cell.font = BODY_FONT
        r += 1

    widths = [20, 26, 10, 16, 13, 6, 8, 10, 7, 11, 11, 11, 7, 7, 9, 12, 6, 15, 16, 18, 15]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def write_techniques(wb, techniques):
    ws = wb.create_sheet('Techniques')
    for i, (_, label) in enumerate(TECH_COLUMNS, start=1):
        ws.cell(row=1, column=i, value=label)
    style_header(ws, len(TECH_COLUMNS))

    r = 2
    for key, row in techniques.items():
        costs = row.get('costs') or {}
        duration = row.get('duration') or {}
        values = {
            'key': key, 'name': row.get('name', ''), 'group': row.get('group', ''),
            'column': row.get('column', ''), 'type': row.get('type', ''),
            'apCost': costs.get('ap'), 'focusCost': costs.get('focus'), 'flowCost': costs.get('flow'),
            'lightningCost': costs.get('lightning'), 'lavaCost': costs.get('lava'), 'range': row.get('range'),
            'durationSequences': duration.get('sequences'), 'targetSide': row.get('targetSide') or '',
            'effect': row.get('effect', ''),
        }
        for i, (field, _) in enumerate(TECH_COLUMNS, start=1):
            cell = ws.cell(row=r, column=i, value=values[field])
            cell.font = BODY_FONT
        r += 1

    widths = [16, 18, 9, 11, 18, 8, 10, 9, 13, 8, 7, 18, 11, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = 'A2'


def write_legend(wb):
    ws = wb.create_sheet('Legend')
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 70
    rows = [
        ('LEGEND — valid values for the enum-like columns', ''),
        ('', ''),
        ('Characters sheet', ''),
        ('key', 'Unique id, lowercase-hyphenated (e.g. "human-bandit"). Adding a new row = a new Character; reusing an existing key overwrites it.'),
        ('category', '"player" or "enemy" — which optgroup it shows up under in-game.'),
        ('species', '"humanoid", "sandman", "crystalBeing", "desertAbomination", or blank (not yet species-tagged — falls back to the universal Death Wounds default and has no Body).'),
        ('woundsMax (y)', 'Personal Down threshold — mutable/levelable, distinct from Species\' Death Wounds (z), which is not spreadsheet-editable (it comes from species alone).'),
        ('Tech: * columns', 'TRUE/FALSE. Blank = TRUE (full access) — matches how the game treats any Character with no techniqueAccess set at all.'),
        ('', ''),
        ('Techniques sheet', ''),
        ('key', 'Unique id, camelCase (e.g. "heartAttack"). Matches the game\'s internal Technique key.'),
        ('group', '"Powers" or "Craft".'),
        ('column', '"Flow", "Focus", "Lightning", or "Lava" — also determines which Technique-List category (Characters sheet) gates it.'),
        ('*Cost columns', 'Leave blank for any resource a Technique doesn\'t spend — only Wound spends Flow with no AP, for example.'),
        ('durationSequences', 'Blank = effect lasts only the Sequence it\'s used in (the default for most Techniques). A number = lasts that many Sequences (Anchor = 3).'),
        ('targetSide', 'Blank = targets Enemies (the default). "ally" = targets the Player, other Players, and their Mini-Party (Embolden, Anchor).'),
        ('', ''),
        ('Pipeline', ''),
        ('To publish changes', 'Save this file, then run: python tools/spreadsheet_to_json.py — regenerates data/characters.json and data/techniques.json, which the game fetches at startup. Refresh the game page to see changes.'),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        ca = ws.cell(row=i, column=1, value=a)
        cb = ws.cell(row=i, column=2, value=b)
        if i == 1:
            ca.font = Font(name='Arial', bold=True, size=13)
        elif b == '' and a != '':
            ca.font = Font(name='Arial', bold=True, size=12)
        else:
            ca.font = BODY_FONT
            cb.font = NOTE_FONT
        cb.alignment = Alignment(wrap_text=True, vertical='top')


def main():
    with open(os.path.join(DATA_DIR, 'characters.json'), encoding='utf-8') as f:
        characters = json.load(f)
    with open(os.path.join(DATA_DIR, 'techniques.json'), encoding='utf-8') as f:
        techniques = json.load(f)

    wb = Workbook()
    write_characters(wb, characters)
    write_techniques(wb, techniques)
    write_legend(wb)
    wb.save(OUT_PATH)
    print('Wrote', OUT_PATH, '—', len(characters), 'Characters,', len(techniques), 'Techniques.')


if __name__ == '__main__':
    main()
