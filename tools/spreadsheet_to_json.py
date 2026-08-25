#!/usr/bin/env python3
"""Regenerates data/*.json from the Character/Technique Templates spreadsheet.

This is the one dDd actually runs: edit data/character-templates.xlsx (add
rows, change stats, flip a Technique-access box), save it, then run this.
The game fetches data/characters.json and data/techniques.json at startup,
so a refresh of the game page picks up whatever's in the spreadsheet — no
code changes, no restart.

Usage: python tools/spreadsheet_to_json.py
Reads:  data/character-templates.xlsx
Writes: data/characters.json, data/techniques.json
"""
import json
import os
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, 'data')
XLSX_PATH = os.path.join(DATA_DIR, 'character-templates.xlsx')

NUMERIC_CHAR_FIELDS = [
    'woundsMax', 'ap', 'flowMax', 'recovery', 'power', 'quickness',
    'resilience', 'lightning', 'lava', 'focus', 'control', 'attunement', 'life',
]
NUMERIC_TECH_FIELDS = ['apCost', 'focusCost', 'flowCost', 'lightningCost', 'lavaCost', 'range', 'durationSequences']

# techniqueAccess went per-Technique (2026-08-24, per dDd — "individual Craft
# and Power, true/false, not all-or-nothing") — one spreadsheet column per
# real Technique key (data/techniques.json), not the 4 old category columns.
TECHNIQUE_ACCESS_KEYS = [
    'wound', 'possess', 'empower', 'resonance', 'current', 'influence',
    'heartAttack', 'skullExplosion', 'message', 'channel',
    'separation', 'infuse', 'bloodkindle', 'seed',
]


def truthy(v, default=True):
    if v is None or v == '':
        return default
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return bool(v)
    return str(v).strip().upper() in ('TRUE', 'YES', 'Y', '1')


def rows_of(ws):
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        if all(v is None or v == '' for v in values):
            continue
        yield dict(zip(headers, values))


def num(v):
    if v is None or v == '':
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _tech_access_label(key):
    # 'heartAttack' -> 'Tech: Heart Attack' — same "Tech: <Title Case>"
    # convention the old 4 category columns used.
    spaced = ''.join(' ' + c if c.isupper() else c for c in key).strip()
    return 'Tech: ' + spaced[:1].upper() + spaced[1:]


TECH_ACCESS_LABEL_TO_KEY = {_tech_access_label(k): k for k in TECHNIQUE_ACCESS_KEYS}


def build_characters(ws):
    label_to_field = {
        'Key': 'key', 'Archetype': 'archetype', 'Category': 'category', 'Species': 'species',
        'Wounds Max (y)': 'woundsMax', 'AP': 'ap', 'Flow': 'flowMax', 'Recovery': 'recovery',
        'Power': 'power', 'Quickness': 'quickness', 'Resilience': 'resilience',
        'Lightning': 'lightning', 'Lava': 'lava', 'Focus': 'focus', 'Control': 'control',
        'Attunement': 'attunement', 'Life': 'life',
    }
    out = {}
    for raw in rows_of(ws):
        row = {label_to_field.get(k, k): v for k, v in raw.items()}
        key = (row.get('key') or '').strip()
        if not key:
            continue
        entry = {
            'archetype': row.get('archetype') or '',
            'category': (row.get('category') or 'character').strip(),
            'species': (row.get('species') or '').strip() or None,
        }
        for field in NUMERIC_CHAR_FIELDS:
            entry[field] = num(row.get(field))
        access = {
            tech_key: truthy(raw.get(label), default=False)
            for label, tech_key in TECH_ACCESS_LABEL_TO_KEY.items()
        }
        if not all(access.values()):
            entry['techniqueAccess'] = access
        out[key] = entry
    return out


def build_techniques(ws):
    label_to_field = {
        'Key': 'key', 'Name': 'name', 'Group': 'group', 'Column': 'column', 'Type': 'type',
        'AP Cost': 'apCost', 'Focus Cost': 'focusCost', 'Flow Cost': 'flowCost',
        'Lightning Cost': 'lightningCost', 'Lava Cost': 'lavaCost', 'Range': 'range',
        'Duration (Sequences)': 'durationSequences', 'Target Side': 'targetSide', 'Effect': 'effect',
    }
    out = {}
    for raw in rows_of(ws):
        row = {label_to_field.get(k, k): v for k, v in raw.items()}
        key = (row.get('key') or '').strip()
        if not key:
            continue
        costs = {}
        for field, resource in [('apCost', 'ap'), ('focusCost', 'focus'), ('flowCost', 'flow'),
                                 ('lightningCost', 'lightning'), ('lavaCost', 'lava')]:
            v = num(row.get(field))
            if v:
                costs[resource] = v
        entry = {
            'name': row.get('name') or '',
            'group': (row.get('group') or '').strip(),
            'column': (row.get('column') or '').strip(),
            'type': row.get('type') or '',
            'costs': costs,
            'range': num(row.get('range')) or 0,
            'effect': row.get('effect') or '',
        }
        duration = num(row.get('durationSequences'))
        if duration:
            entry['duration'] = {'sequences': duration}
        target_side = (row.get('targetSide') or '').strip()
        if target_side:
            entry['targetSide'] = target_side
        out[key] = entry
    return out


def main():
    wb = load_workbook(XLSX_PATH, data_only=True)
    characters = build_characters(wb['Characters'])
    techniques = build_techniques(wb['Techniques'])

    with open(os.path.join(DATA_DIR, 'characters.json'), 'w', encoding='utf-8') as f:
        json.dump(characters, f, indent=2, ensure_ascii=False)
        f.write('\n')
    with open(os.path.join(DATA_DIR, 'techniques.json'), 'w', encoding='utf-8') as f:
        json.dump(techniques, f, indent=2, ensure_ascii=False)
        f.write('\n')

    print('Wrote characters.json (' + str(len(characters)) + ') and techniques.json (' + str(len(techniques)) + ')')


if __name__ == '__main__':
    main()
